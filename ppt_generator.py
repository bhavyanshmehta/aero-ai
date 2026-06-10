import os
from pptx import Presentation
from pptx.util import Inches, Pt

def set_text_frame(shape, text, font_size=Pt(18)):
    if not hasattr(shape, "text_frame") or shape.text_frame is None:
        return
    shape.text_frame.clear()
    p = shape.text_frame.paragraphs[0]
    p.text = text
    for run in p.runs:
        run.font.size = font_size

def create_aero_ai_ppt():
    sample_path = "/Users/bhavyansh/Documents/ANSHU/EMS_Presentation.pptx"
    output_path = "/Users/bhavyansh/Documents/ANSHU/Aero_Ai_Presentation.pptx"
    mockup_path = "/Users/bhavyansh/.gemini/antigravity/brain/78619008-c279-4455-ab60-54afbd6bf248/aero_ai_ui_mockup_1778694684116.png"
    arch_path = "/Users/bhavyansh/.gemini/antigravity/brain/c7b3ed5d-7421-4f35-ba12-006eeda9cc9b/scratch/report_media/image2.png"

    if not os.path.exists(sample_path):
        print("Sample PPT not found!")
        return

    prs = Presentation(sample_path)
    
    # Slide 1: Title
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if hasattr(shape, "text") and shape.text:
            if "Education Management System" in shape.text:
                shape.text = "AERO AI: A MULTIMODAL INTELLIGENT CHAT ASSISTANT"
            if "Presented By" in shape.text:
                shape.text = "Presented By\nBhavyansh Mehta (Reg No. 24BCONXXXX)"
            elif any(name in shape.text for name in ["Piyush", "Ayushi", "Aarya", "Priyanshu", "Rohan"]):
                shape.text = ""

    # Slide 3: Problem Statement
    slide3 = prs.slides[2]
    content3 = (
        "• Session Volatility: Standard AI interfaces lack persistent session management.\n"
        "• Limited Multimodality: Existing tools fail to provide native image-based reasoning.\n"
        "• UI/UX Gap: A need for premium, local workspaces with modern aesthetics."
    )
    for shape in slide3.shapes:
        if hasattr(shape, "text_frame") and shape.text_frame:
            if len(shape.text) > 10:
                set_text_frame(shape, content3)

    # Slide 4: Introduction
    slide4 = prs.slides[3]
    content4 = (
        "• Aero Ai is a next-generation multimodal chat platform powered by Gemini 1.5 Flash.\n"
        "• Designed for high-performance text and image analysis in a unified interface.\n"
        "• Bridges the gap between complex AI models and professional-grade user interfaces."
    )
    for shape in slide4.shapes:
        if hasattr(shape, "text") and "PROJECT OVERVIEW" in shape.text:
            shape.text = content4

    # Slide 6: Roles (Diagram)
    slide6 = prs.slides[5]
    if len(slide6.shapes) > 9:
        try:
            slide6.shapes[7].text = "End User"
            slide6.shapes[8].text = "FastAPI Backend"
            slide6.shapes[9].text = "Gemini AI Engine"
        except: pass

    # Slide 7: Tools
    slide7 = prs.slides[6]
    content7 = (
        "• Backend: Python & FastAPI (Asynchronous Performance)\n"
        "• Frontend: HTML5, CSS3, Vanilla JavaScript (Glassmorphism)\n"
        "• Database: SQLite3 (Local Persistence)\n"
        "• AI SDK: Google Generative AI (Gemini 1.5 Flash)"
    )
    for shape in slide7.shapes:
        if hasattr(shape, "text_frame") and shape.text_frame:
            if len(shape.text) > 20:
                set_text_frame(shape, content7)

    # Slide 9: Module Description
    slide9 = prs.slides[8]
    content9 = (
        "• AI Integration: Gemini 1.5 Flash handles multimodal queries asynchronously.\n"
        "• Persistence Layer: SQLite manages chat sessions and history retrieval.\n"
        "• UI Engine: Modern Glassmorphism design with responsive sidebar navigation.\n"
        "• Technical: Code highlighting using Highlight.js and Marked.js integration."
    )
    for shape in slide9.shapes:
        if hasattr(shape, "text_frame") and shape.text_frame:
            if len(shape.text) > 30:
                set_text_frame(shape, content9)

    # Slide 10: Code
    slide10 = prs.slides[9]
    code = """
@app.post("/chat")
async def chat(message: str, image: UploadFile = None):
    # Gemini Multimodal Logic
    model = genai.GenerativeModel('gemini-1.5-flash')
    img_data = await image.read() if image else None
    response = model.generate_content([message, img_data] if img_data else message)
    return {"response": response.text}
    """
    for shape in slide10.shapes:
        if hasattr(shape, "text") and "login" in shape.text:
            set_text_frame(shape, code, font_size=Pt(14))

    # Slide 12: Conclusion
    slide12 = prs.slides[11]
    content12 = (
        "• Aero Ai integrates Gemini AI with FastAPI for a premium chat experience.\n"
        "• All functional objectives met, including persistence and multimodal support.\n"
        "• Future: Multi-user auth, voice-to-text, and local model integration."
    )
    for shape in slide12.shapes:
        if hasattr(shape, "text_frame") and shape.text_frame:
            if len(shape.text) > 30:
                set_text_frame(shape, content12)

    # Slide 13: Screenshots
    slide13 = prs.slides[12]
    if os.path.exists(mockup_path):
        slide13.shapes.add_picture(mockup_path, Inches(0.5), Inches(1.5), width=Inches(9))

    prs.save(output_path)
    print(f"Presentation saved to {output_path}")

if __name__ == "__main__":
    create_aero_ai_ppt()
